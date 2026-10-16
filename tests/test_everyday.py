"""What the everyday-work scenarios taught: each test pins one fix found by simulating agents at work."""

import os
import sys
import time
import zipfile

import pytest

from minion.agent import a11y, docs
from minion.agent.a11y import Element, compact, format_elements
from minion.agent.perceive import Perception, _changed_lines, _clear, _overlaps_picture, _same_content
from minion.agent.server import Agent, busy_text, parse_step, run_captured

# The desktop's own agent runs on Linux: tests that run its shell commands or resolve its paths need Linux.
linux_only = pytest.mark.skipif(not sys.platform.startswith("linux"), reason="the desktop agent runs on Linux")


def el(role, name, depth=0, y=0, x=100, w=80, h=18, **kw):
    return Element(0, role, name, depth, x, y, w, h, **kw)


def web(role, name, depth=0, y=0, **kw):
    return el(role, name, depth, y, in_doc=True, **kw)


# -- running commands ---------------------------------------------------------------------------
@linux_only
def test_run_returns_when_a_background_job_keeps_the_pipe_open():
    # `(cd dir && server &)`: the subshell that runs the job holds stdout open as long as the server runs
    t0 = time.monotonic()
    res = run_captured(["bash", "-c", "(cd /tmp && sleep 20 &) ; echo started"], "/tmp", 10, None)
    assert time.monotonic() - t0 < 3
    assert res["stdout"] == "started\n" and res["code"] == 0 and not res["timed_out"]


@linux_only
def test_run_keeps_only_the_ends_of_huge_output():
    # `find -L` in a symlink loop printed gigabytes and the desktop ran out of memory
    res = run_captured(["bash", "-c", "head -c 20000000 /dev/zero | tr '\\0' a"], "/tmp", 30, None)
    assert "bytes omitted" in res["stdout"] and len(res["stdout"]) < 70_000


@linux_only
def test_run_still_times_out_and_passes_input():
    res = run_captured(["bash", "-c", "sleep 5"], "/tmp", 0.5, None)
    assert res["timed_out"]
    res = run_captured(["cat"], "/tmp", 5, b"x" * 200_000)
    assert len(res["stdout"]) > 60_000  # (trimmed to the output limit, but all of it went through)


# -- the screen as text ---------------------------------------------------------------------------
def test_labels_that_only_repeat_their_control_are_dropped():
    els = [
        web("label", "Free Free", 0, 0),
        web("radio button", "Free", 1, 0, states=["unchecked"]),
        web("label", "Phone +972 50 555 0101", 0, 30),
        web("entry", "Phone", 1, 30, value="+972 50 555 0101"),
        web("label", "Remember me on this device", 0, 60),
        web("check box", "Remember me", 1, 60),
    ]
    lines = format_elements(compact(els))
    assert lines[0].startswith('@1 radio button "Free"')
    assert lines[1].startswith('@2 entry "Phone"')
    assert lines[2].startswith('@3 label "Remember me on this device"')  # it says more than the control


def test_wrappers_in_wrappers_go_too():
    els = [web("section", "Small", 0, 0), web("label", "Small", 1, 0), web("radio button", "Small", 2, 0)]
    lines = format_elements(compact(els))
    assert len(lines) == 1 and lines[0].startswith('@1 radio button "Small"')


def test_web_tables_come_one_row_per_line_with_their_controls_inline():
    els = [web("heading", "People", 0, 0)]
    for c, name in enumerate(["ID", "Name"]):
        els += [
            el("column header", name, 0, 40, x=100 + 150 * c, in_doc=True),
            el("link", name, 1, 40, x=105 + 150 * c, in_doc=True),
        ]
    for r in range(2):
        for c, value in enumerate([str(1001 + r), ["Kim", "Sam"][r]]):
            els.append(el("table cell", value, 0, 80 + 34 * r, x=100 + 150 * c, in_doc=True))
        els.append(el("table cell", "", 0, 80 + 34 * r, x=400, in_doc=True))
        els.append(el("push button", "Approve", 1, 80 + 34 * r, x=405, in_doc=True))
    out = compact(els)
    lines = format_elements(out)
    assert lines[1].startswith('@2 table "3 rows x 3 columns"')
    assert lines[2].strip().startswith("| @") and 'link "Name"' in lines[2]
    assert lines[3].strip().startswith("| 1001 | Kim | @") and 'push button "Approve"' in lines[3]
    assert len(lines) == 5  # heading, table, 3 rows: no line per cell
    cells = [e for e in out if e.inline]
    assert cells and all(e.id for e in cells)  # still there for @ids and text targets


def test_a_toolbar_is_one_line_and_blank_names_are_dropped():
    els = [el("tool bar", "Standard", 0, 0, w=1000)]
    els += [
        el("push button", n, 1, 0, x=10 + 30 * k, states=s)
        for k, (n, s) in enumerate(
            [("Open", []), ("Save", []), ("Undo", ["disabled"]), ("Bold", ["checked"]), ("Italic", ["unchecked"])]
        )
    ]
    lines = format_elements(compact(els))
    assert (
        lines[1]
        .strip()
        .startswith('buttons: @2 "Open" · @3 "Save" · @4 "Undo" [disabled] · @5 "Bold" [checked] · @6 "Italic"')
    )
    assert "unchecked" not in lines[1]


def test_unnamed_check_boxes_are_named_by_the_text_next_to_them():
    els = []
    for k, t in enumerate(["milk", "bread"]):
        els += [
            web("list item", "", 0, 30 * k),
            web("check box", "", 1, 30 * k, states=["unchecked"]),
            web("label", t, 1, 30 * k),
        ]
    lines = format_elements(compact(els))
    assert lines[0].startswith('@1 check box "next to milk" [unchecked]')


def test_controls_that_share_a_label_say_where_they_are():
    els = [web("heading", "Mailbox", 0, 0)]
    for k, box in enumerate(["Inbox", "Spam"]):
        els += [web("heading", box, 0, 40 + 60 * k), web("push button", "Empty", 0, 60 + 60 * k)]
    els += [web("section", "USB desk lamp … Add to cart", 0, 200, own=True), web("paragraph", "$6.75", 1, 210)]
    els += [web("push button", "Add to cart", 1, 220), web("section", "HDMI cable … Add to cart", 0, 260, own=True)]
    els += [web("push button", "Add to cart", 1, 270)]
    lines = format_elements(compact(els))
    assert any(line.endswith('in "Inbox"') and '"Empty"' in line for line in lines)
    assert any(line.endswith('in "Spam"') and '"Empty"' in line for line in lines)
    assert any(line.endswith('in "USB desk lamp"') for line in lines)


def test_a_page_dialog_goes_first_and_the_page_keeps_its_depth():
    page = [web("heading", "Account settings", 6, 100), web("push button", "Change display name", 6, 140)]
    dialog = [web("dialog", "The page at http://localhost says:", 5, 400, modal=True)]
    dialog += [
        web("entry", "Enter your new display name:", 6, 420, modal=True),
        web("push button", "OK", 6, 460, modal=True),
    ]
    out, _ = a11y.browser_compact(page + dialog)
    assert [e.role for e in out[:3]] == ["dialog", "entry", "push button"]
    assert out[0].depth == 0 and out[1].depth == 1
    assert next(e for e in out if e.name == "Account settings").depth == 0


def test_only_a_real_picture_gets_the_no_text_hint():
    photo = el("image", "photo.jpg", 0, 100, x=100, w=400, h=300)
    assert _overlaps_picture((100, 100, 400, 300), [photo])
    assert not _overlaps_picture((100, 500, 400, 200), [photo])


def test_ocr_in_a_well_described_window_keeps_only_clear_words():
    class Line:
        def __init__(self, text, conf):
            self.text, self.conf = text, conf

    assert _clear(Line("Total due", 88))
    assert not _clear(Line("oO. Bp,", 80))
    assert not _clear(Line("/ io 2 8", 90))
    assert not _clear(Line("Rows processed", 50))


# -- changes between looks -------------------------------------------------------------------------
def test_new_text_shows_up_in_the_changes_not_just_its_box():
    p = Perception(None, None, (1280, 800))
    section = web("section", "", 0, 300)
    section.lines = ["Welcome aboard, Maya!", "Account created."]
    section.id = 7
    old = Perception._view({}, None, [], [])
    new = Perception._view({}, None, [section], [])
    text = p._diff(old, new)["text"]
    assert "Welcome aboard, Maya!" in text and "Account created." in text


def test_a_terminal_that_printed_more_counts_as_changed():
    a = ("terminal", "", "$ python3", "", ())
    b = ("terminal", "", "$ python3\n>>> exit()\n$", "", ())
    assert not _same_content(a, b)
    assert _same_content(("paragraph", "", "Long text…", "", ()), ("paragraph", "", "Long text that goes on", "", ()))
    assert _changed_lines(("terminal", ("$ ls",)), ("terminal", ("$ ls", "a.txt", "$"))) == ["a.txt", "$"]


def test_full_view_reasons_say_what_happened():
    old = {"focused": "w1", "windows": {"w1": ("Save As", True)}, "items": {}, "time": 0}
    assert Perception._diff(old, {"focused": None, "windows": {}, "items": {}, "time": 1}) == {
        "full": "no window is open now"
    }
    closed = Perception._diff(old, {"focused": "w2", "windows": {"w2": ("Notepad", True)}, "items": {}, "time": 1})
    assert closed == {"full": "the window you looked at last is closed"}


def test_web_elements_never_inherit_an_id_by_place():
    p = Perception(None, None, (1280, 800))
    first = web("push button", "Reserve a seat", 0, 300)
    first.node = object()
    p._assign([first])
    other = web("push button", "Reserve a seat", 0, 300)  # after a scroll: another session's button, same spot
    other.node = object()
    p._assign([other])
    assert other.id != first.id
    field = el("entry", "Find", 0, 50)
    field.node, field.text = object(), ""
    p._assign([field])
    again = el("entry", "Find", 0, 50)  # LibreOffice recreated the field; it now holds "14%"
    again.node, again.text = object(), "14%"
    p._assign([again])
    assert again.id == field.id


# -- actions -------------------------------------------------------------------------------------------
def test_act_strings_for_the_new_steps():
    assert parse_step('wait_for "Report ready" 60') == ("wait_for", {"text": "Report ready", "timeout": 60.0})
    assert parse_step("wait_for Generating --gone") == ("wait_for", {"text": "Generating", "gone": True})
    assert parse_step("wait_for Page 2") == ("wait_for", {"text": "Page 2"})
    assert parse_step("scroll_to Entry 87") == ("scroll_to", {"text": "Entry 87"})
    assert parse_step('find "Sapiens"') == ("scroll_to", {"text": "Sapiens"})
    assert parse_step("rclick @12") == ("click", {"button": "right", "target": "@12"})
    assert parse_step("dclick 'My file'") == ("click", {"count": 2, "target": "My file"})
    assert parse_step("close Notepad") == ("close", {"window": "Notepad"})
    assert parse_step("copy @121") == ("copy", {"target": "@121"})


def test_busy_texts_are_status_lines_not_titles():
    assert busy_text("section", "Generating... this can take a while.")
    assert busy_text("status", "Saving")
    assert busy_text("paragraph", "Loading more…")
    # statuses that can stay for good are not "busy": an order's "Processing", a task's "In progress", a promo
    assert busy_text("paragraph", "Processing") is None
    assert busy_text("section", "In progress") is None
    assert busy_text("paragraph", "Saving 20%") is None
    assert busy_text("heading", "Uploading files") is None
    assert busy_text("paragraph", "Loading docks are on the north side of the building.") is None
    assert busy_text("push button", "Download") is None


def test_the_reply_says_which_files_an_action_saved_or_downloaded(tmp_path):
    home = str(tmp_path)
    os.makedirs(f"{home}/Documents/letters")
    os.makedirs(f"{home}/Downloads")
    os.makedirs(f"{home}/.cache")
    (tmp_path / "Documents/report.docx").write_text("x")
    (tmp_path / "Downloads/old.bin").write_text("y" * 5000)
    agent = object.__new__(Agent)
    agent.home = home
    before = agent._files_state()
    time.sleep(0.01)
    with zipfile.ZipFile(tmp_path / "Documents/report.docx", "w") as z:  # (a real save: a Word document)
        z.writestr("word/document.xml", "<w:document/>")
    size = os.path.getsize(tmp_path / "Documents/report.docx")
    (tmp_path / "Documents/letters/Avi.pdf").write_text("z" * 3000)
    (tmp_path / "Downloads/data.csv").write_text("")
    (tmp_path / "Downloads/data.csv.part").write_text("partial")
    (tmp_path / ".cache/noise").write_text("n")
    os.remove(tmp_path / "Downloads/old.bin")
    notes = agent._files_changed(before, agent._files_state())
    assert sorted(notes) == sorted(
        [
            "downloading ~/Downloads/data.csv (not finished yet)",
            "new file ~/Documents/letters/Avi.pdf (3 KB)",
            f"saved ~/Documents/report.docx ({size} bytes, Word format)",
            "deleted ~/Downloads/old.bin",
        ]
    )  # (the order follows modification times, which can tie on some systems)


# -- documents -----------------------------------------------------------------------------------------
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
CT = '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'


def make_zip(path, parts: dict) -> str:
    with zipfile.ZipFile(path, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return str(path)


def test_docx_shows_its_table_of_contents_changes_comments_and_page(tmp_path):
    body = (
        f"<w:document {W}><w:body>"
        "<w:sdt><w:sdtContent><w:p><w:r><w:t>Table of Contents</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Results 1</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        '<w:p><w:ins w:id="1" w:author="Ada"><w:r><w:t>15%</w:t></w:r></w:ins>'
        '<w:del w:id="2" w:author="Ada"><w:r><w:delText>14%</w:delText></w:r></w:del></w:p>'
        "<w:p><w:r><w:drawing/></w:r></w:p>"
        '<w:sectPr><w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/>'
        '<w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134"/></w:sectPr>'
        "</w:body></w:document>"
    )
    comments = f'<w:comments {W}><w:comment w:id="0" w:author="Ada"><w:p><w:r><w:t>Check with finance</w:t></w:r></w:p></w:comment></w:comments>'
    path = make_zip(
        tmp_path / "r.docx", {"[Content_Types].xml": CT, "word/document.xml": body, "word/comments.xml": comments}
    )
    text, _ = docs.read(path)
    assert "Table of Contents" in text and "Results 1" in text
    assert "15%" in text and "14%" not in text.split("---")[0]
    assert "[image]" in text
    assert "Tracked changes not yet accepted: 1 insertions, 1 deletions" in text
    assert "Comment by Ada: Check with finance" in text
    assert "Page: 29.7 x 21.0 cm (landscape), margins 2.00 / 2.00 / 2.00 / 2.00 cm" in text


def test_xlsx_shows_bold_frozen_panes_conditional_formats_and_charts(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    for row in [("Region", "Revenue"), ("North", 120000), ("South", 80000)]:
        ws.append(row)
    ws["A1"].font = ws["B1"].font = Font(bold=True)
    ws.freeze_panes = "A2"
    red = PatternFill(bgColor="FF0000", fill_type="solid")
    ws.conditional_formatting.add("B2:B3", CellIsRule(operator="greaterThan", formula=["100000"], fill=red))
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(chart, "D2")
    wb.save(tmp_path / "s.xlsx")
    text, _ = docs.read(str(tmp_path / "s.xlsx"))
    assert "Bold: A1, B1" in text
    assert "Frozen panes: rows and columns before A2 stay in view" in text
    assert "Conditional format: B2:B3 greaterThan 100000 → fill #FF0000" in text
    assert "Chart: column, clustered: 'Sales'!B1 = 'Sales'!$B$2:$B$3" in text


def test_a_picture_without_text_is_described_not_misread(tmp_path, monkeypatch):
    from PIL import Image

    from minion.agent import ocr

    Image.new("RGB", (640, 360), (40, 90, 200)).save(tmp_path / "frame.png")
    monkeypatch.setattr(docs.shutil, "which", lambda tool: "/usr/bin/" + tool)
    monkeypatch.setattr(ocr, "ensure", lambda langs: (["eng"], ""))
    monkeypatch.setattr(ocr, "page_text", lambda img, langs: ("y", 1))
    assert docs._ocr(tmp_path / "frame.png").startswith("[PNG image, 640x360: no text in it that OCR could read")


def test_convert_takes_a_file_name_and_a_web_page_converts_as_a_document(tmp_path, monkeypatch):
    (tmp_path / "s.csv").write_text("a,b\n1,2\n")
    out = docs.convert(str(tmp_path / "s.csv"), "sorted.md")
    assert out == str(tmp_path / "sorted.md") and (tmp_path / "sorted.md").exists()
    out = docs.convert(str(tmp_path / "s.csv"), str(tmp_path / "out" / "x.md"))
    assert (tmp_path / "out" / "x.md").exists()
    seen = {}

    def fake_run(cmd, **kw):
        seen["cmd"] = cmd
        outdir = cmd[cmd.index("--outdir") + 1]
        open(os.path.join(outdir, "page.docx"), "w").close()

        class R:
            returncode, stdout, stderr = 0, b"", b""

        return R()

    (tmp_path / "page.html").write_text("<h1>Hi</h1>")
    monkeypatch.setattr(docs, "_soffice", lambda: "soffice")
    monkeypatch.setattr(docs.subprocess, "run", fake_run)
    docs.convert(str(tmp_path / "page.html"), "docx")
    assert "--infilter=HTML (StarWriter)" in seen["cmd"]


# -- the command line ----------------------------------------------------------------------------------
def test_cli_open_takes_paths_the_shell_expanded_and_several_urls():
    from minion.cli import URLISH, desktop_path

    home = os.path.expanduser("~")
    if home not in ("/", "/home/agent"):
        assert desktop_path(home + "/Pictures/photo.png") == "~/Pictures/photo.png"
    assert URLISH.match("https://wikipedia.org") and URLISH.match("python.org")
    assert not URLISH.match("gimp")


# -- second round (from the run on the fixed Minion) ------------------------------------------------
def test_spreadsheet_cells_come_one_row_per_line():
    els = [Element(1, "sheet", "Sheet1", 0, 0, 100, 1000, 500, text="visible A1:B2")]
    k = 2
    for r, row in enumerate([["Region", "Revenue"], ["North", "120000"]], 1):
        for c, v in enumerate(row):
            els.append(Element(k, "cell", f"{'AB'[c]}{r}", 1, 100 + 80 * c, 100 + 20 * r, 80, 20, text=v))
            k += 1
    els.append(Element(k, "cell", "B3", 1, 180, 160, 80, 20, text="120000", value="=B2"))
    lines = format_elements(els)
    assert lines[1].strip() == "row 1: @2 A1 Region · @3 B1 Revenue"
    assert lines[3].strip() == "row 3: @6 B3 120000 (=B2)"


def test_a_label_beside_its_control_is_dropped():
    els = [web("radio button", "Small", 0, 20, states=["unchecked"]), web("label", "Small", 0, 22, text="…")]
    els += [web("section", "Small", 1, 22), web("radio button", "Large", 0, 50, states=["checked"])]
    lines = format_elements(compact(els))
    assert [line.split(" (")[0] for line in lines] == [
        '@1 radio button "Small" [unchecked]',
        '@2 radio button "Large" [checked]',
    ]


def test_numbers_misread_with_lookalike_letters_are_fixed():
    from minion.agent.ocr import _digits

    assert _digits("7i") == "71" and _digits("2O26") == "2026" and _digits("10:3O") == "10:30"
    assert _digits("Oil") == "Oil" and _digits("hello") == "hello" and _digits("I") == "I"


def test_menu_items_named_by_their_command_still_match():
    class Node:
        def __init__(self, name, kids=()):
            self.name, self.kids = name, list(kids)

    menu = Node("Track Changes", [Node("record"), Node("show"), Node("manage...")])
    calls = {"_children": lambda n, limit=200: n.kids, "_role": lambda n: "menu item", "_name": lambda n: n.name}
    saved = {k: getattr(a11y, k) for k in calls}
    try:
        for k, f in calls.items():
            setattr(a11y, k, f)
        found, names = a11y._menu_child(menu, "Record Changes")
    finally:
        for k, f in saved.items():
            setattr(a11y, k, f)
    assert found is menu.kids[0] and names == ["record", "show", "manage"]


@linux_only
def test_open_takes_a_profile_app_and_a_file_windows_paths_included():
    agent = object.__new__(Agent)
    agent.apps = [{"name": "Notepad", "id": "notepad", "command": "wine notepad"}]
    agent.home = "/home/agent"
    cmd, kind = agent._resolve_open(r"notepad C:\users\agent\Documents\notes.txt")
    assert kind == "app" and cmd == "wine notepad 'C:\\users\\agent\\Documents\\notes.txt'"


def test_date_fields_take_a_date_in_any_clear_form():
    import datetime as dt

    for text in ("2027-03-15", "15 March 2027", "March 15, 2027", "15.03.2027"):
        day = None
        for fmt in Agent.DATE_FORMATS:
            try:
                day = dt.datetime.strptime(text.replace(",", ""), fmt).date()
                break
            except ValueError:
                continue
        assert day == dt.date(2027, 3, 15), text


def test_the_changes_view_keeps_where_a_repeated_control_is():
    els = [web("heading", "Claim 2046", 0, 10), web("push button", "Approve", 0, 30)]
    els += [web("heading", "Claim 2047", 0, 60), web("push button", "Approve", 0, 80)]
    for k, e in enumerate(els, 1):
        e.id = k
    view = Perception._view({}, None, els, [])
    assert view["items"][4][2].endswith('in "Claim 2047"')


# -- from the code review ----------------------------------------------------------------------------
def test_windows_paths_keep_their_backslashes_everywhere():
    from minion.agent.server import _escapes as desk
    from minion.tools import _escapes as host

    for text, typed in [
        (r"C:\users\agent\note.txt", "C:\\users\\agent\\note.txt"),
        (r"C:\users\agent\note.txt\n", "C:\\users\\agent\\note.txt\n"),
        (r"C:\temp", "C:\\temp"),
        (r"X:\t5\tY:\t7\n", "X:\t5\tY:\t7\n"),  # a spreadsheet row, not a path
        (r"C:\\temp\\new", "C:\\temp\\new"),  # the escape convention still works
        (r"hello\nworld", "hello\nworld"),
    ]:
        assert desk(text) == typed and host(text) == typed, text


def test_a_label_with_a_link_in_it_stays():
    els = [web("check box", "I agree to the Terms", 0, 20, states=["unchecked"])]
    els += [web("label", "I agree to the Terms", 0, 22), web("link", "Terms", 1, 22)]
    lines = format_elements(compact(els))
    assert any('link "Terms"' in line for line in lines)


def test_menu_command_names_match_only_command_style_items():
    class Node:
        def __init__(self, name):
            self.name = name

    file_menu = [Node("Close"), Node("Export..."), Node("Save")]
    calls = {"_children": lambda n, limit=200: n, "_role": lambda n: "menu item", "_name": lambda n: n.name}
    saved = {k: getattr(a11y, k) for k in calls}
    try:
        for k, f in calls.items():
            setattr(a11y, k, f)
        assert a11y._menu_child(file_menu, "Close All")[0] is None  # not "Close"
        assert a11y._menu_child(file_menu, "Export to PDF")[0] is None  # not "Export..."
    finally:
        for k, f in saved.items():
            setattr(a11y, k, f)


def test_web_ids_are_never_refound_by_place():
    p = Perception(None, None, (1280, 800))
    gone = web("push button", "Delete", 0, 300)
    assert p._refind(gone) is None  # the next row's "Delete" may sit there now


@linux_only
def test_open_keeps_an_apps_arguments_apart(tmp_path):
    agent = object.__new__(Agent)
    agent.apps = [{"name": "Calc", "id": "calc", "command": "libreoffice --calc"}]
    agent.home = str(tmp_path)
    (tmp_path / "a b.csv").write_text("x")
    cmd, _ = agent._resolve_open(f"calc {tmp_path}/a b.csv")
    assert cmd == f"libreoffice --calc '{tmp_path}/a b.csv'"  # one file, spaces and all
    cmd, _ = agent._resolve_open("calc --nologo /tmp/x.csv")
    assert cmd == "libreoffice --calc --nologo /tmp/x.csv"  # a flag and a file: two arguments


def test_a_pipe_is_not_read_as_a_one():
    from minion.agent.ocr import _digits

    assert _digits("|") == "|" and _digits("||") == "||" and _digits("7i") == "71"
